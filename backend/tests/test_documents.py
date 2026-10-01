from __future__ import annotations

import unittest
from io import BytesIO

from docx import Document

from app.documents.parser import normalize_document_text, parse_document
from app.documents.service import parse_document_upload


class DocumentParserTest(unittest.TestCase):
    def test_accepts_short_notes_and_rejects_whitespace_only(self) -> None:
        self.assertEqual(parse_document("note.md", "按主题整理资料。".encode()), "按主题整理资料。")
        with self.assertRaises(ValueError):
            parse_document("note.txt", b" \n\t")

    def test_parses_utf8_text_resume(self) -> None:
        text = parse_document("resume.txt", "张三\nPython 开发工程师\n五年后端开发经验".encode())
        self.assertIn("Python 开发工程师", text)

    def test_parse_response_does_not_infer_retired_career_fields(self) -> None:
        result = parse_document_upload(
            "resume.txt",
            "姓名：张三\n求职意向：后端工程师\n期望城市：上海\n技能：Python、Docker".encode(),
            "fast",
        )

        self.assertNotIn("suggested_profile", result)
        self.assertNotIn("suggested_skills", result)

    def test_parse_response_preserves_original_and_returns_separate_redacted_text(self) -> None:
        result = parse_document_upload(
            "resume.txt",
            (
                "李明\n"
                "求职方向：AI 应用研发 电话：13812345678\n"
                "邮箱：candidate@example.com\n"
                "技能：Python\n"
            ).encode(),
            "fast",
        )

        self.assertIn("求职方向：AI 应用研发", result["text"])
        self.assertIn("技能：Python", result["text"])
        self.assertIn("李明", result["text"])
        self.assertIn("13812345678", result["text"])
        self.assertIn("candidate@example.com", result["text"])
        self.assertNotIn("13812345678", result["redacted_text"])
        self.assertNotIn("candidate@example.com", result["redacted_text"])
        self.assertIn("已隐藏", result["redacted_text"])

    def test_parses_docx_paragraphs_and_tables(self) -> None:
        document = Document()
        document.add_heading("个人简历", level=1)
        document.add_paragraph("负责 FastAPI 和 Agent 平台开发")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "技能"
        table.cell(0, 1).text = "Python"
        stream = BytesIO()
        document.save(stream)

        text = parse_document("resume.docx", stream.getvalue())

        self.assertIn("Agent 平台开发", text)
        self.assertIn("技能\tPython", text)

    def test_normalizes_private_use_bullets_and_invisible_characters(self) -> None:
        text = parse_document(
            "resume.txt",
            "智能会议总结\n\uf0b7 基于 LangChain 搭建统一网关\u200b\n".encode(),
        )

        self.assertIn("- 基于 LangChain 搭建统一网关", text)
        self.assertNotIn("\uf0b7", text)
        self.assertNotIn("\u200b", text)

    def test_unwraps_pdf_visual_line_breaks_without_merging_headings(self) -> None:
        text = normalize_document_text(
            "个人优势\n"
            "- 熟悉 Python / FastAPI / Docker 等技术栈，\n"
            "擅长实时语音链路与多模型协同。\n"
            "- 独立完成业务内容生产链路，有效提升业务内容产出效\n"
            "率 60%+\n"
            "工作经历\n"
            "某公司\n"
            "AI 应用工程师\n"
        )

        self.assertIn("Docker 等技术栈，擅长实时语音链路与多模型协同。", text)
        self.assertIn("产出效率 60%+", text)
        self.assertIn("\n工作经历\n", text)
        self.assertIn("\n某公司\n", text)
        self.assertIn("AI 应用工程师", text)
        self.assertNotIn("某公司AI 应用工程师", text)

    def test_splits_jammed_contact_and_certificate_fields(self) -> None:
        text = normalize_document_text(
            "小程\n"
            "邮箱: [邮箱已隐藏] 英语: CET-6GitHub: https://github.com/Xiaolang-d3\n"
            "求职意向：后端工程师\n"
        )

        self.assertIn("邮箱: [邮箱已隐藏]", text)
        self.assertIn("英语: CET-6", text)
        self.assertIn("GitHub: https://github.com/Xiaolang-d3", text)
        self.assertNotIn("CET-6GitHub", text)
        self.assertIn("\n求职意向：后端工程师", text)

    def test_rejects_unsupported_or_empty_files(self) -> None:
        with self.assertRaises(ValueError):
            parse_document("resume.pages", b"not supported")
        with self.assertRaises(ValueError):
            parse_document("resume.txt", b"")


if __name__ == "__main__":
    unittest.main()
